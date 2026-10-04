"""Manual CSV / XLSX uploads, validated row by row against a strict schema. One bad row rejects the whole file."""
from __future__ import annotations

import csv
import io
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator

from .models import Kind, Txn, to_minor


class FileRow(BaseModel):
    """Expected columns. Unknown columns are an error, so a shifted or renamed header cannot slip through."""
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)

    external_id: str = Field(min_length=1, max_length=200)
    kind: Kind
    amount: Decimal                       # major units, e.g. 1234.56; negative for refunds / chargebacks
    fee: Decimal = Decimal(0)
    currency: str = Field(pattern=r"^[A-Z]{3}$")
    occurred_at: datetime                 # ISO 8601; a bare date means midnight UTC
    reference: str | None = None
    payout_id: str | None = None
    customer: str | None = None

    @field_validator("amount", "fee", mode="before")
    @classmethod
    def _decimal(cls, v: Any) -> Any:
        if isinstance(v, float):          # Excel hands numbers over as floats: drop binary noise (1e-6 is far
            v = repr(round(v, 6))         # below a cent), but real extra decimals still fail in to_minor
        try:
            d = Decimal(str(v))
        except InvalidOperation as exc:
            raise ValueError("not a number") from exc
        if not d.is_finite():
            raise ValueError("not a finite number")
        return d

    @field_validator("reference", "payout_id", "customer", mode="before")
    @classmethod
    def _blank_is_none(cls, v: Any) -> Any:
        return None if v is None or str(v).strip() == "" else str(v)

    @field_validator("external_id", "kind", "currency", mode="before")
    @classmethod
    def _text(cls, v: Any) -> Any:
        return "" if v is None else str(v)


class FileRejected(ValueError):
    def __init__(self, errors: list[dict[str, Any]]):
        super().__init__(f"{len(errors)} invalid row(s)")
        self.errors = errors


def _rows(filename: str, content: bytes) -> list[tuple[int, dict[str, Any]]]:
    """(row number as the person sees it in the file, {column: value})."""
    if filename.lower().endswith(".xlsx"):
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        assert sheet is not None
        values = list(sheet.iter_rows(values_only=True))
        header = [str(h).strip() if h is not None else "" for h in values[0]] if values else []
        return [(n, dict(zip(header, row))) for n, row in enumerate(values[1:], start=2) if any(c is not None for c in row)]
    if filename.lower().endswith(".csv"):
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        return [(reader.line_num, row) for row in reader]
    raise FileRejected([{"row": 0, "error": "only .csv and .xlsx files are accepted"}])


def parse(filename: str, content: bytes, source: str = "file") -> list[Txn]:
    errors: list[dict[str, Any]] = []
    txns: list[Txn] = []
    for n, raw in _rows(filename, content):
        try:
            row = FileRow.model_validate(raw)
            when = row.occurred_at if row.occurred_at.tzinfo else row.occurred_at.replace(tzinfo=UTC)
            txns.append(Txn(source=source, external_id=row.external_id, kind=row.kind,
                            amount_minor=to_minor(row.amount, row.currency), fee_minor=to_minor(row.fee, row.currency),
                            currency=row.currency, occurred_at=when, reference=row.reference,
                            payout_id=row.payout_id, customer=row.customer))
        except ValidationError as exc:
            errors += [{"row": n, "column": ".".join(map(str, e["loc"])), "error": e["msg"]} for e in exc.errors()]
        except ValueError as exc:
            errors.append({"row": n, "error": str(exc)})
    if not txns and not errors:
        errors.append({"row": 0, "error": "the file has no data rows"})
    if errors:
        raise FileRejected(errors)
    return txns
